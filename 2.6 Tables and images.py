# pip install Pillow library beforehand

from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl import load_workbook

wb = load_workbook('Pie.xlsx')
ws = wb.active

tab = Table(displayName='Table1', ref='A1:B5')
style = TableStyleInfo(name='TableStyleMedium9', showFirstColumn=False, showLastColumn=False,
                       showRowStripes=True, showColumnStripes=True)

tab.tableStyleInfo = style
# add the style to the table

ws.add_table(tab)
# add table to the worksheet

wb.save('table.xlsx')

img = Image('madecraft.jpg')
# create an image variable and assign the image to it, if the image is not in the cwd, path needs to be assigned

img.height = img.height*0.25
img.width = img.width*0.25
# if the current img is too big, adjust the image size by multiplying 0.25 

ws.add_image(img, 'C8')
#  add image to the worksheet

wb.save('image.xlsx')
